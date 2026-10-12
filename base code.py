# ============================================================
# Traditional Programming vs Machine Learning
# Loan Approval Decision using Excel Input / Output
# ============================================================

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter
from collections import Counter
import math

# ===================== FILE NAMES =====================
INPUT_EXCEL = "Loan_Input.xlsx"      # change if needed
OUTPUT_EXCEL = "Loan_Input.xlsx"


# ===================== HELPER =====================
def safe_number(value):
    """Convert value to float. Returns None if blank/invalid."""
    if value is None or (isinstance(value, float) and math.isnan(value)):
        return None
    if isinstance(value, str) and value.strip() == "":
        return None
    try:
        return float(value)
    except (ValueError, TypeError):
        return None


# ===================== 1. TRADITIONAL PROGRAMMING =====================
def traditional_decision(income, credit_score, existing_loan):
    """
    Rules written by the programmer:
    Income >= 50000 AND Credit Score >= 700 AND Existing Loan <= 20000
    --> APPROVED, otherwise REJECTED
    """
    if income is None or credit_score is None or existing_loan is None:
        return "INVALID DATA"

    if income >= 50000 and credit_score >= 700 and existing_loan <= 20000:
        return "APPROVED"
    return "REJECTED"


# ===================== 2. MACHINE LEARNING (Decision Tree) =====================
class DecisionNode:
    def __init__(self, feature=None, threshold=None, left=None, right=None, value=None):
        self.feature = feature
        self.threshold = threshold
        self.left = left
        self.right = right
        self.value = value


def entropy(y):
    if len(y) == 0:
        return 0
    counts = Counter(y)
    total = len(y)
    return -sum((c / total) * math.log2(c / total) for c in counts.values() if c > 0)


def information_gain(X_col, y, threshold):
    left_mask = X_col <= threshold
    right_mask = ~left_mask
    if left_mask.sum() == 0 or right_mask.sum() == 0:
        return 0.0
    parent_ent = entropy(y)
    n = len(y)
    left_ent = entropy(y[left_mask])
    right_ent = entropy(y[right_mask])
    child_ent = (left_mask.sum() / n) * left_ent + (right_mask.sum() / n) * right_ent
    return parent_ent - child_ent


def best_split(X, y, feature_names):
    best_gain = -1
    best_feature = None
    best_threshold = None
    for feat in range(X.shape[1]):
        values = sorted(set(X[:, feat]))
        for i in range(len(values) - 1):
            thr = (values[i] + values[i + 1]) / 2
            gain = information_gain(X[:, feat], y, thr)
            if gain > best_gain:
                best_gain = gain
                best_feature = feat
                best_threshold = thr
    return best_feature, best_threshold, best_gain


def build_tree(X, y, feature_names, depth=0, max_depth=3):
    if len(set(y)) == 1:
        return DecisionNode(value=y[0])
    if depth >= max_depth or len(y) < 2:
        majority = Counter(y).most_common(1)[0][0]
        return DecisionNode(value=majority)

    feat, thr, gain = best_split(X, y, feature_names)
    if feat is None or gain <= 0:
        majority = Counter(y).most_common(1)[0][0]
        return DecisionNode(value=majority)

    left_mask = X[:, feat] <= thr
    right_mask = ~left_mask
    left_child = build_tree(X[left_mask], y[left_mask], feature_names, depth + 1, max_depth)
    right_child = build_tree(X[right_mask], y[right_mask], feature_names, depth + 1, max_depth)
    return DecisionNode(feature=feat, threshold=thr, left=left_child, right=right_child)


def predict_one(node, x):
    if node.value is not None:
        return node.value
    if x[node.feature] <= node.threshold:
        return predict_one(node.left, x)
    return predict_one(node.right, x)


def export_tree_text(node, feature_names, indent=""):
    if node.value is not None:
        label = "APPROVED" if node.value == 1 else "REJECTED"
        return f"{indent}-> {label}\n"
    name = feature_names[node.feature]
    text = f"{indent}if {name} <= {node.threshold:.1f}:\n"
    text += export_tree_text(node.left, feature_names, indent + "  ")
    text += f"{indent}else:  # {name} > {node.threshold:.1f}\n"
    text += export_tree_text(node.right, feature_names, indent + "  ")
    return text


def train_ml_model(training_df):
    """Train Decision Tree on historical data"""
    feature_names = ["Income", "Credit_Score", "Existing_Loan"]
    X = training_df[feature_names].values.astype(float)
    y = (training_df["Actual_Decision"] == "APPROVED").astype(int).values
    tree = build_tree(X, y, feature_names, max_depth=3)
    return tree, feature_names


def ml_decision(tree, feature_names, income, credit_score, existing_loan):
    if income is None or credit_score is None or existing_loan is None:
        return "INVALID DATA"
    x = [income, credit_score, existing_loan]
    pred = predict_one(tree, x)
    return "APPROVED" if pred == 1 else "REJECTED"


# ===================== 3. MAIN =====================
def main():
    print("=" * 65)
    print("Traditional Programming  vs  Machine Learning")
    print("Loan Approval Decision - Excel Input / Output")
    print("=" * 65)

    # Read Excel
    training_df = pd.read_excel(INPUT_EXCEL, sheet_name="Training_Data", header=2)
    customers_df = pd.read_excel(INPUT_EXCEL, sheet_name="Customers", header=2)

    # Clean data
    training_df = training_df.dropna(subset=["Income", "Credit_Score", "Existing_Loan"]).copy()
    customers_df = customers_df.dropna(subset=["Customer_ID"]).copy()
    customers_df = customers_df[customers_df["Customer_ID"].astype(str).str.startswith("C")].copy()

    print(f"\nTraining samples loaded : {len(training_df)}")
    print(f"Customers to evaluate  : {len(customers_df)}")

    # Train ML model from historical data
    tree, feature_names = train_ml_model(training_df)
    print("\nDecision Tree trained successfully (from historical data).")
    print("\nRules learned by the ML model:")
    print(export_tree_text(tree, feature_names))

    # Apply both methods
    traditional_results = []
    ml_results = []
    match_results = []

    for _, row in customers_df.iterrows():
        income = safe_number(row["Income"])
        credit = safe_number(row["Credit_Score"])
        loan = safe_number(row["Existing_Loan"])

        trad = traditional_decision(income, credit, loan)
        ml = ml_decision(tree, feature_names, income, credit, loan)

        traditional_results.append(trad)
        ml_results.append(ml)
        match_results.append("SAME" if trad == ml else "DIFFERENT")

    customers_df["Traditional_Decision"] = traditional_results
    customers_df["ML_Decision"] = ml_results
    customers_df["Match"] = match_results

    print("\n--- Decisions ---")
    cols = ["Customer_ID", "Name", "Income", "Credit_Score", "Existing_Loan",
            "Traditional_Decision", "ML_Decision", "Match"]
    print(customers_df[cols].to_string(index=False))

    # Write results back to Excel
    wb = load_workbook(INPUT_EXCEL)

    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(bold=True, color="FFFFFF")
    thin = Border(left=Side(style='thin'), right=Side(style='thin'),
                  top=Side(style='thin'), bottom=Side(style='thin'))
    center = Alignment(horizontal="center", vertical="center")
    green_fill = PatternFill("solid", fgColor="C6EFCE")
    green_font = Font(color="006100", bold=True)
    red_fill = PatternFill("solid", fgColor="FFC7CE")
    red_font = Font(color="9C0006", bold=True)
    yellow_fill = PatternFill("solid", fgColor="FFEB9C")
    yellow_font = Font(color="9C5700", bold=True)
    gray_fill = PatternFill("solid", fgColor="D9D9D9")
    gray_font = Font(color="595959", bold=True)

    ws = wb["Customers"]

    result_headers = ["Traditional_Decision", "ML_Decision", "Match"]
    for col_idx, h in enumerate(result_headers, 6):
        cell = ws.cell(row=3, column=col_idx, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center
        cell.border = thin

    for row_idx, (_, row) in enumerate(customers_df.iterrows(), 4):
        for col, val in [(6, row["Traditional_Decision"]),
                         (7, row["ML_Decision"]),
                         (8, row["Match"])]:
            cell = ws.cell(row=row_idx, column=col, value=val)
            cell.border = thin
            cell.alignment = center

            if val == "APPROVED":
                cell.fill = green_fill
                cell.font = green_font
            elif val == "REJECTED":
                cell.fill = red_fill
                cell.font = red_font
            elif val == "INVALID DATA":
                cell.fill = gray_fill
                cell.font = gray_font
            elif val == "SAME":
                cell.fill = green_fill
                cell.font = green_font
            else:
                cell.fill = yellow_fill
                cell.font = yellow_font

    for col in range(6, 9):
        ws.column_dimensions[get_column_letter(col)].width = 20

    wb.save(OUTPUT_EXCEL)
    print(f"\nResults written back to: {OUTPUT_EXCEL}")
    print("=" * 65)


if __name__ == "__main__":
    main()
